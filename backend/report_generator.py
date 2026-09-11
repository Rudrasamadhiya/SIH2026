# backend/report_generator.py
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
from datetime import datetime
import json

def generate_clinical_report(patient_info: dict, conversation_history: str, ai_summary: dict) -> bytes:
    """
    Generates a structured medical report in .docx format.
    Returns the file as bytes (for direct download).
    """
    doc = Document()
    
    # --- Styles ---
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # --- Header ---
    header = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.text = "MEDICAL KIOSK - CLINICAL SUMMARY REPORT"
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_para.runs[0].font.size = Pt(10)
    header_para.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    # --- Title ---
    title = doc.add_heading('Patient Consultation Summary', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- Patient Demographics ---
    doc.add_heading('1. Patient Demographics', level=1)
    table = doc.add_table(rows=4, cols=2, style='Light Grid Accent 1')
    data = [
        ("Name / ID", patient_info.get("name", "Anonymous (Kiosk User)")),
        ("Age / Gender", patient_info.get("age_gender", "Not Provided")),
        ("Date / Time", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Kiosk ID", patient_info.get("kiosk_id", "KIOSK-001"))
    ]
    for i, (k, v) in enumerate(data):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v

    # --- Chief Complaint ---
    doc.add_heading('2. Chief Complaint', level=1)
    doc.add_paragraph(ai_summary.get("chief_complaint", "Not explicitly stated."))

    # --- History of Present Illness (SOCRATES) ---
    doc.add_heading('3. History of Present Illness (SOCRATES Framework)', level=1)
    socrates = ai_summary.get("socrates_analysis", {})
    if socrates:
        s_table = doc.add_table(rows=len(socrates)+1, cols=2, style='Light Grid Accent 1')
        s_table.rows[0].cells[0].text = "Parameter"
        s_table.rows[0].cells[1].text = "Details"
        for idx, (key, val) in enumerate(socrates.items()):
            s_table.rows[idx+1].cells[0].text = key.capitalize()
            s_table.rows[idx+1].cells[1].text = val
    else:
        doc.add_paragraph("SOCRATES analysis not available.")

    # --- AYUSH Observations ---
    doc.add_heading('4. AYUSH / Lifestyle Observations', level=1)
    ayush = ai_summary.get("ayush_notes", {})
    if ayush:
        for k, v in ayush.items():
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{k.capitalize()}: ").bold = True
            p.add_run(v)
    else:
        doc.add_paragraph("No specific AYUSH parameters recorded.")

    # --- AI Clinical Impression ---
    doc.add_heading('5. AI Clinical Impression & Red Flags', level=1)
    impression = ai_summary.get("impression", "No impression generated.")
    p = doc.add_paragraph()
    run = p.add_run("⚠️ " if ai_summary.get("emergency_flag") else "✅ ")
    run.font.color.rgb = RGBColor(255, 0, 0) if ai_summary.get("emergency_flag") else RGBColor(0, 128, 0)
    run.bold = True
    p.add_run(impression)

    # --- Full Conversation Log (Appendix) ---
    doc.add_heading('Appendix: Full Conversation Transcript', level=1)
    doc.add_paragraph(conversation_history)

    # --- Save to Bytes ---
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()